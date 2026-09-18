import os
import logging
from typing import List, Dict, Any, Optional
from app.core.config import settings

logger = logging.getLogger("document-parser")

class DocumentParser:
    """
    Parses PDF, DOCX, and TXT documents.
    Implements ADR-04 (Markdown table preservation via pdfplumber)
    and ADR-05 (Multimodal vision chart extraction via Gemini 2.5 Flash).
    """

    @staticmethod
    def parse(file_path: str) -> List[Dict[str, Any]]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found at path: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            return DocumentParser._parse_pdf(file_path)
        elif ext == ".docx":
            return DocumentParser._parse_docx(file_path)
        elif ext in [".txt", ".md", ".csv"]:
            return DocumentParser._parse_text(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    @staticmethod
    def _parse_pdf(file_path: str) -> List[Dict[str, Any]]:
        pages = []
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for idx, page in enumerate(pdf.pages):
                    page_num = idx + 1
                    raw_text = page.extract_text() or ""
                    page_content_parts = [raw_text.strip()] if raw_text.strip() else []

                    # ADR-04: Detect and convert tabular data to Markdown tables
                    try:
                        tables = page.extract_tables()
                        for table in tables:
                            md_table = DocumentParser._format_markdown_table(table)
                            if md_table:
                                page_content_parts.append(f"\n[Table on Page {page_num}]:\n{md_table}")
                    except Exception as te:
                        logger.debug(f"Table extraction skipped on page {page_num}: {te}")

                    # ADR-05: Check for figures/charts and extract with Gemini 2.5 Flash Vision
                    if settings.GEMINI_API_KEY and page.images:
                        try:
                            chart_analysis = DocumentParser._analyze_visual_charts(page, page_num)
                            if chart_analysis and chart_analysis.strip() != "NONE":
                                page_content_parts.append(
                                    f"\n[Visual Chart & Figure Data - Page {page_num}]:\n{chart_analysis}"
                                )
                        except Exception as ve:
                            logger.debug(f"Vision analysis skipped on page {page_num}: {ve}")

                    combined_text = "\n\n".join(page_content_parts).strip()
                    if combined_text:
                        pages.append({
                            "page_number": page_num,
                            "text": combined_text
                        })

        except ImportError:
            logger.warning("pdfplumber not available; falling back to basic pypdf extraction.")
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            for idx, page in enumerate(reader.pages):
                text = (page.extract_text() or "").strip()
                if text:
                    pages.append({"page_number": idx + 1, "text": text})

        return pages

    @staticmethod
    def _format_markdown_table(table: List[List[Optional[str]]]) -> Optional[str]:
        if not table or len(table) < 2:
            return None

        # Clean cells
        cleaned_table = []
        for row in table:
            cleaned_row = [str(cell).strip().replace("\n", " ") if cell is not None else "" for cell in row]
            if any(cleaned_row):
                cleaned_table.append(cleaned_row)

        if len(cleaned_table) < 2:
            return None

        col_count = max(len(row) for row in cleaned_table)
        # Pad shorter rows
        for row in cleaned_table:
            while len(row) < col_count:
                row.append("")

        header = "| " + " | ".join(cleaned_table[0]) + " |"
        separator = "| " + " | ".join(["---"] * col_count) + " |"
        data_rows = ["| " + " | ".join(row) + " |" for row in cleaned_table[1:]]

        return "\n".join([header, separator] + data_rows)

    @staticmethod
    def _analyze_visual_charts(page: Any, page_num: int) -> Optional[str]:
        """ADR-05: Renders page image to extract chart metrics via Gemini 2.5 Flash."""
        try:
            import io
            from google import genai
            from google.genai import types

            # Render page to image at 150 DPI
            pil_image = page.to_image(resolution=150).original
            img_byte_arr = io.BytesIO()
            pil_image.save(img_byte_arr, format="PNG")
            img_bytes = img_byte_arr.getvalue()

            client = genai.Client(api_key=settings.GEMINI_API_KEY)
            prompt = (
                f"Examine this page (page {page_num}). If there are any charts, graphs, or visual data figures, "
                "extract the numeric data into a clear Markdown table and describe the key trend in 1-2 sentences. "
                "If there are no charts or graphs (only standard text, logos, or icons), respond with exactly 'NONE'."
            )

            response = client.models.generate_content(
                model=settings.VISION_MODEL or "gemini-2.5-flash",
                contents=[
                    types.Part.from_bytes(data=img_bytes, mime_type="image/png"),
                    prompt
                ]
            )
            return response.text.strip() if response.text else None
        except Exception as e:
            logger.debug(f"Gemini chart extraction failed: {e}")
            return None

    @staticmethod
    def _parse_docx(file_path: str) -> List[Dict[str, Any]]:
        try:
            import docx
            doc = docx.Document(file_path)
            parts = []

            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())

            # Extract tables in docx
            for table in doc.tables:
                rows = []
                for row in table.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
                md_table = DocumentParser._format_markdown_table(rows)
                if md_table:
                    parts.append(f"\n[Table]:\n{md_table}")

            full_text = "\n\n".join(parts)
            return [{"page_number": 1, "text": full_text}]
        except ImportError:
            return [{"page_number": 1, "text": "DOCX reader not installed"}]

    @staticmethod
    def _parse_text(file_path: str) -> List[Dict[str, Any]]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            full_text = f.read().strip()
        return [{"page_number": 1, "text": full_text}]
